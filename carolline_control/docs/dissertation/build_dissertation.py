"""Assemble MECH5845M dissertation Word document from markdown chapters."""



from __future__ import annotations



import re

import shutil

from pathlib import Path



from docx import Document

from docx.enum.text import WD_LINE_SPACING

from docx.oxml import OxmlElement

from docx.oxml.ns import qn

from docx.shared import Mm, Pt, Inches



REPO = Path(__file__).resolve().parents[3]

DISS = REPO / "carolline_control" / "docs" / "dissertation"

FIG = REPO / "carolline_control" / "plots" / "report_figures"

HYBRID_LOG = REPO / "carolline_control" / "logs" / "hybrid_maze"

OUT = REPO / "dissertation" / "CAROLLINE_MSc_Report.docx"



TITLE = (

    "Simulation and Control of a Caged Hybrid Rolling–Flying Aerial Vehicle "

    "with Autonomous Navigation and SLAM in MuJoCo"

)



BODY_FONT = "Arial"

BODY_SIZE = Pt(11)

CAPTION_SIZE = Pt(10)



FIG_MAP = {

    "Figure 3.1": "fig_architecture.png",

    "Figure 3.2": "fig_esc_characterization.png",

    "Figure 4.1": "fig_mode_fsm.png",

    "Figure 5.1": "fig_actual_vs_desired.png",

    "Figure 5.2": "fig_slam_vs_ground_truth.png",

    "Figure 5.3": "hybrid_maze_trajectory_xy.png",

    "Figure 5.4": "hybrid_maze_altitude.png",

    "Figure 5.5": "hybrid_maze_mode_timeline.png",

    "Figure 5.6": "hybrid_maze_actual_vs_desired_xyz.png",

}



HYBRID_FIG_SOURCES = [

    "hybrid_maze_trajectory_xy.png",

    "hybrid_maze_altitude.png",

    "hybrid_maze_mode_timeline.png",

    "hybrid_maze_actual_vs_desired_xyz.png",

]





def _set_margins(doc: Document) -> None:

    for section in doc.sections:

        section.top_margin = Mm(25)

        section.bottom_margin = Mm(25)

        section.right_margin = Mm(25)

        section.left_margin = Mm(38)





def _set_default_font(doc: Document) -> None:

    style = doc.styles["Normal"]

    style.font.name = BODY_FONT

    style.font.size = BODY_SIZE

    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE





def _style_body(paragraph, *, single_spacing: bool = False) -> None:

    if single_spacing:

        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    else:

        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    for run in paragraph.runs:

        run.font.name = BODY_FONT

        run.font.size = BODY_SIZE





def _style_caption(paragraph) -> None:

    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    for run in paragraph.runs:

        run.font.name = BODY_FONT

        run.font.size = CAPTION_SIZE





def _style_reference(paragraph) -> None:

    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    for run in paragraph.runs:

        run.font.name = BODY_FONT

        run.font.size = CAPTION_SIZE





def _add_heading(doc: Document, text: str, level: int) -> None:

    p = doc.add_heading(text, level=level)

    for run in p.runs:

        run.font.name = BODY_FONT

        run.font.size = BODY_SIZE if level > 1 else Pt(12)

    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    p.paragraph_format.space_before = Pt(12 if level == 1 else 6)

    p.paragraph_format.space_after = Pt(6)





def _add_body(doc: Document, text: str) -> None:

    text = text.strip()

    if not text:

        return

    p = doc.add_paragraph(text)

    _style_body(p)





def _add_page_number_field(paragraph, *, roman: bool = False) -> None:

    run = paragraph.add_run()

    fld_begin = OxmlElement("w:fldChar")

    fld_begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")

    instr.set(qn("xml:space"), "preserve")

    instr.text = " PAGE \\* ROMAN \\* MERGEFORMAT " if roman else " PAGE \\* MERGEFORMAT "

    fld_sep = OxmlElement("w:fldChar")

    fld_sep.set(qn("w:fldCharType"), "separate")

    fld_end = OxmlElement("w:fldChar")

    fld_end.set(qn("w:fldCharType"), "end")

    run._r.extend([fld_begin, instr, fld_sep, fld_end])

    run.font.name = BODY_FONT

    run.font.size = CAPTION_SIZE





def _configure_section_numbering(section, *, fmt: str | None = None, start: int | None = None) -> None:

    sect_pr = section._sectPr

    for child in list(sect_pr):

        if child.tag == qn("w:pgNumType"):

            sect_pr.remove(child)

    if fmt is not None or start is not None:

        pg_num = OxmlElement("w:pgNumType")

        if fmt:

            pg_num.set(qn("w:fmt"), fmt)

        if start is not None:

            pg_num.set(qn("w:start"), str(start))

        sect_pr.append(pg_num)





def _clear_footer(section) -> None:
    section.footer.is_linked_to_previous = False
    footer = section.footer
    for paragraph in footer.paragraphs:
        paragraph.clear()
    if section.different_first_page_header_footer:
        first_footer = section.first_page_footer
        first_footer.is_linked_to_previous = False
        for paragraph in first_footer.paragraphs:
            paragraph.clear()


def _add_section_footer_page_number(section, *, roman: bool = False) -> None:
    section.footer.is_linked_to_previous = False
    footer = section.footer
    for paragraph in footer.paragraphs:
        paragraph.clear()
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    paragraph.alignment = 1  # center
    _add_page_number_field(paragraph, roman=roman)


def _start_new_section(doc: Document, *, roman: bool = False, arabic_restart: bool = False) -> None:
    doc.add_section()
    section = doc.sections[-1]
    section.footer.is_linked_to_previous = False
    _set_margins(doc)
    if roman:
        _configure_section_numbering(section, fmt="upperRoman", start=2)
        _add_section_footer_page_number(section, roman=True)
    elif arabic_restart:
        _configure_section_numbering(section, fmt="decimal", start=1)
        _add_section_footer_page_number(section, roman=False)
    else:
        section.different_first_page_header_footer = True
        _configure_section_numbering(section, fmt="upperRoman", start=1)





def _add_table_from_md(doc: Document, lines: list[str], caption: str | None = None) -> None:

    rows = [ln for ln in lines if ln.strip().startswith("|")]

    if len(rows) < 2:

        return

    headers = [c.strip() for c in rows[0].strip("|").split("|")]

    body = []

    for row in rows[2:]:

        body.append([c.strip() for c in row.strip("|").split("|")])

    table = doc.add_table(rows=1 + len(body), cols=len(headers))

    table.style = "Table Grid"

    for j, h in enumerate(headers):

        cell = table.rows[0].cells[j]

        cell.text = h

        for p in cell.paragraphs:

            _style_body(p, single_spacing=True)

    for i, row in enumerate(body, start=1):

        for j, val in enumerate(row):

            cell = table.rows[i].cells[j]

            cell.text = val

            for p in cell.paragraphs:

                _style_body(p, single_spacing=True)

    if caption:

        cap = doc.add_paragraph(caption)

        _style_caption(cap)





def _parse_md_file(path: Path) -> list[tuple[str, str]]:

    """Return list of (kind, content) where kind is heading|body|table|figure."""

    if not path.exists():

        return []

    items: list[tuple[str, str]] = []

    lines = path.read_text(encoding="utf-8").splitlines()

    buf: list[str] = []

    in_table = False



    def flush_body() -> None:

        nonlocal buf

        if buf:

            items.append(("body", "\n".join(buf).strip()))

            buf = []



    for line in lines:

        if line.startswith("#"):

            flush_body()

            level = len(line) - len(line.lstrip("#"))

            title = line.lstrip("#").strip()

            items.append(("heading", f"{level}|{title}"))

            continue

        if line.strip().startswith("|"):

            if not in_table:

                flush_body()

                in_table = True

                buf = [line]

            else:

                buf.append(line)

            continue

        if in_table:

            items.append(("table", "\n".join(buf)))

            buf = []

            in_table = False

        for fig_label, fname in FIG_MAP.items():

            if fig_label in line:

                items.append(("figure", f"{fig_label}|{fname}|{line.strip()}"))

        if line.strip():

            buf.append(line.strip())

        elif buf:

            flush_body()

    flush_body()

    if in_table and buf:

        items.append(("table", "\n".join(buf)))

    return items





def _add_cover(doc: Document) -> None:

    section = doc.sections[0]

    section.different_first_page_header_footer = True
    _configure_section_numbering(section, fmt="upperRoman", start=1)
    _clear_footer(section)



    for _ in range(6):

        doc.add_paragraph()



    p = doc.add_paragraph("MECH5845M Professional Engineering Project")

    p.runs[0].bold = True

    p.runs[0].font.size = Pt(14)

    p.runs[0].font.name = BODY_FONT

    doc.add_paragraph()

    t = doc.add_paragraph(TITLE)

    t.runs[0].bold = True

    t.runs[0].font.size = Pt(14)

    t.runs[0].font.name = BODY_FONT

    doc.add_paragraph()

    doc.add_paragraph()



    block = doc.add_paragraph()

    block.alignment = 2  # right

    for i, line in enumerate(

        [

            "Author: ######### (Name and SID)",

            "Supervisor: #########",

            "Examiner: #########",

            "Industrial Mentor: N/A",

            "Date of Submission: #########",

        ]

    ):

        if i:

            block.add_run("\n")

        run = block.add_run(line)

        run.font.name = BODY_FONT

        run.font.size = BODY_SIZE



    doc.add_page_break()





def _add_declaration(doc: Document) -> None:

    _add_heading(doc, "School of Mechanical Engineering", 1)

    doc.add_paragraph("MECH5845M Professional Engineering Project")

    doc.add_paragraph()

    p = doc.add_paragraph(

        "This project report presents my own work and does not contain any "

        "unacknowledged work from any other sources."

    )

    _style_body(p)

    doc.add_paragraph()

    p2 = doc.add_paragraph("Signed: _________________________    Date: ______________")

    _style_body(p2)

    doc.add_page_break()





def _ensure_figures() -> None:

    FIG.mkdir(parents=True, exist_ok=True)

    for fname in HYBRID_FIG_SOURCES:

        src = HYBRID_LOG / fname

        dst = FIG / fname

        if src.is_file():

            shutil.copy2(src, dst)

    required = set(FIG_MAP.values())

    missing = [f for f in required if not (FIG / f).is_file()]

    if missing:

        import subprocess

        import sys



        gen_script = FIG / "generate_report_figures.py"

        if gen_script.is_file():

            subprocess.run([sys.executable, str(gen_script)], cwd=str(REPO), check=False)

        for fname in HYBRID_FIG_SOURCES:

            src = HYBRID_LOG / fname

            dst = FIG / fname

            if src.is_file():

                shutil.copy2(src, dst)

        try:

            from carolline_control.navigation.hybrid_maze_report import (

                regenerate_hybrid_maze_report_from_dir,

            )



            regenerate_hybrid_maze_report_from_dir(HYBRID_LOG)

            for fname in HYBRID_FIG_SOURCES:

                src = HYBRID_LOG / fname

                dst = FIG / fname

                if src.is_file():

                    shutil.copy2(src, dst)

        except Exception:

            pass

        render_script = REPO / "carolline_control" / "scripts" / "render_maze_slam_map.py"

        if render_script.is_file():

            subprocess.run([sys.executable, str(render_script)], cwd=str(REPO), check=False)

        traj_script = REPO / "carolline_control" / "plots" / "trajectory_tracking.py"

        log_csv = REPO / "carolline_control" / "logs" / "flight_log.csv"

        if traj_script.is_file() and log_csv.is_file():

            subprocess.run(

                [

                    sys.executable,

                    str(traj_script),

                    "--log",

                    str(log_csv),

                    "--output",

                    str(FIG / "fig_actual_vs_desired.png"),

                ],

                cwd=str(REPO),

                check=False,

            )





def _render_chapter_items(doc: Document, items: list[tuple[str, str]]) -> None:

    pending_table_caption: str | None = None

    for kind, content in items:

        if kind == "heading":

            level_s, title = content.split("|", 1)

            _add_heading(doc, title, int(level_s))

        elif kind == "body":

            for para in content.split("\n\n"):

                if para.startswith("```"):

                    continue

                if para.startswith("|"):

                    _add_table_from_md(doc, para.splitlines(), caption=pending_table_caption)

                    pending_table_caption = None

                elif para.startswith("**Table") or para.startswith("Table "):

                    pending_table_caption = re.sub(r"\*\*(.+?)\*\*", r"\1", para)

                else:

                    clean = re.sub(r"\*\*(.+?)\*\*", r"\1", para)

                    clean = re.sub(r"`(.+?)`", r"\1", clean)

                    _add_body(doc, clean)

        elif kind == "table":

            _add_table_from_md(doc, content.splitlines(), caption=pending_table_caption)

            pending_table_caption = None

        elif kind == "figure":

            label, fname, caption = content.split("|", 2)

            fig_path = FIG / fname

            if fig_path.exists():

                doc.add_picture(str(fig_path), width=Inches(5.5))

                cap_text = caption if caption.startswith(label) else f"{label}: {caption}"

                cap = doc.add_paragraph(cap_text)

                if cap.runs:

                    cap.runs[0].italic = True

                _style_caption(cap)





def build() -> Path:

    _ensure_figures()

    doc = Document()

    _set_margins(doc)

    _set_default_font(doc)

    _add_cover(doc)



    _start_new_section(doc, roman=True)

    _add_declaration(doc)



    _add_heading(doc, "Abstract", 1)

    for kind, content in _parse_md_file(DISS / "abstract.md"):

        if kind == "body":

            _add_body(doc, content)

    doc.add_page_break()



    _add_heading(doc, "Contents", 1)

    p = doc.add_paragraph(

        "[Insert Word Table of Contents field here: References → Table of Contents → Update before printing]"

    )

    _style_body(p)

    doc.add_page_break()



    _start_new_section(doc, arabic_restart=True)



    chapter_files = [

        "ch01_introduction.md",

        "ch02_literature_review.md",

        "ch03_modelling.md",

        "ch04_control_design.md",

        "ch05_navigation_results.md",

        "ch06_conclusion.md",

    ]



    for i, fname in enumerate(chapter_files):

        if i > 0:

            doc.add_page_break()

        items = _parse_md_file(DISS / fname)

        _render_chapter_items(doc, items)



    doc.add_page_break()

    _add_heading(doc, "References", 1)

    for kind, content in _parse_md_file(DISS / "references.md"):

        if kind == "body":

            for para in content.split("\n\n"):

                p = doc.add_paragraph(para)

                _style_reference(p)



    doc.add_page_break()

    _add_heading(doc, "Appendices", 1)

    for kind, content in _parse_md_file(DISS / "appendices.md"):

        if kind == "heading":

            _, title = content.split("|", 1)

            _add_heading(doc, title, 2)

        elif kind == "body":

            for para in content.split("\n\n"):

                if para.startswith("```"):

                    for code_line in para.strip("`").splitlines():

                        if code_line.startswith("powershell"):

                            continue

                        _add_body(doc, code_line)

                else:

                    _add_body(doc, para)

        elif kind == "table":

            _add_table_from_md(doc, content.splitlines())



    OUT.parent.mkdir(parents=True, exist_ok=True)
    try:
        doc.save(OUT)
        return OUT
    except PermissionError:
        fallback = OUT.with_name("CAROLLINE_MSc_Report_build.docx")
        doc.save(fallback)
        print(
            f"WARNING: Could not overwrite {OUT} (file may be open in Word). "
            f"Saved to {fallback} instead."
        )
        return fallback





if __name__ == "__main__":

    path = build()

    print(f"Dissertation written to: {path}")


